#!/usr/bin/env python3
"""DC Express DCIM 사용자 매뉴얼 PowerPoint 생성 스크립트"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# === Color Palette (Light Theme) ===
SLIDE_BG = RGBColor(0xFF, 0xFF, 0xFF)       # 흰색 배경
ACCENT_BLUE = RGBColor(0x00, 0x70, 0xC0)    # 메인 파란색 (진하게)
ACCENT_CYAN = RGBColor(0x00, 0xA5, 0x85)    # 포인트 시안 (진하게)
ACCENT_PURPLE = RGBColor(0x6C, 0x2E, 0xD6)  # 보라색
ACCENT_ORANGE = RGBColor(0xE0, 0x8A, 0x00)  # 주황색 (진하게)
ACCENT_RED = RGBColor(0xDC, 0x2E, 0x2E)     # 빨간색
ACCENT_GREEN = RGBColor(0x16, 0xA3, 0x4A)   # 초록색
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_PRIMARY = RGBColor(0x1F, 0x1F, 0x1F)   # 본문 텍스트 (거의 검정)
TEXT_SECONDARY = RGBColor(0x5A, 0x5A, 0x6E) # 보조 텍스트 (짙은 회색)
TEXT_MUTED = RGBColor(0x8C, 0x8C, 0x9A)     # 서브불릿 (중간 회색)
LIGHT_GRAY = RGBColor(0xE5, 0xE7, 0xEB)     # 경계선
CARD_BG = RGBColor(0xF8, 0xF9, 0xFB)        # 카드 배경 (아주 연한 회색)
CARD_BORDER = RGBColor(0xE0, 0xE0, 0xE8)    # 카드 테두리
SECTION_BG_TINT = RGBColor(0xF0, 0xF4, 0xF8) # 섹션 슬라이드 배경
SECTION_COLORS = [
    ACCENT_BLUE,    # 1. Dashboard
    ACCENT_CYAN,    # 2. Server Monitoring
    ACCENT_PURPLE,  # 3. Rack View
    RGBColor(0x06, 0xB6, 0xD4),  # 4. Digital Twin
    ACCENT_ORANGE,  # 5. BMC/Redfish
    RGBColor(0x84, 0xCC, 0x16),  # 6. K8s
    ACCENT_GREEN,   # 7. Workloads
    RGBColor(0xEC, 0x48, 0x99),  # 8. Infrastructure
    ACCENT_RED,     # 9. Alerts
    RGBColor(0x8B, 0x5C, 0xF6),  # 10. Settings
    RGBColor(0x14, 0xB8, 0xA6),  # 11. Capacity/Reports
    RGBColor(0x6B, 0x72, 0x80),  # 12. Others
]

current_section_idx = 0

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_slide_bg(slide, color=SLIDE_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    line = shape.line
    if border_color:
        line.color.rgb = border_color
        line.width = border_width
    else:
        line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_accent_bar(slide, left, top, width, height, color=ACCENT_BLUE):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_text_box(slide, left, top, width, height, text, font_size=14, color=TEXT_PRIMARY,
                 bold=False, alignment=PP_ALIGN.LEFT, font_name="맑은 고딕"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_text(slide, left, top, width, height, bullets, font_size=14, color=TEXT_PRIMARY,
                    spacing=Pt(6), font_name="맑은 고딕", line_spacing=1.3):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        is_sub = bullet.startswith("•")
        text = bullet.lstrip("• ").strip()

        if is_sub:
            p.text = text
            p.level = 1
            p.font.size = Pt(font_size - 1)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(2)
            p.space_after = Pt(2)
            p.font.name = font_name
        else:
            p.text = text
            p.level = 0
            p.font.size = Pt(font_size)
            p.font.color.rgb = color
            p.font.bold = False
            p.space_before = spacing
            p.space_after = Pt(2)
            p.font.name = font_name

        p.line_spacing = line_spacing

    return txBox


def make_cover_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    # gradient-like top bar
    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), ACCENT_BLUE)
    add_accent_bar(slide, Inches(0), Inches(0.06), SLIDE_W, Inches(0.03), ACCENT_CYAN)

    # center content
    add_text_box(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(1),
                 "DC Express", font_size=52, color=ACCENT_BLUE, bold=True)
    add_text_box(slide, Inches(1.5), Inches(2.8), Inches(10), Inches(0.8),
                 "Data Center Infrastructure Management System", font_size=24, color=TEXT_SECONDARY)

    add_accent_bar(slide, Inches(1.5), Inches(3.8), Inches(3), Inches(0.04), ACCENT_CYAN)

    add_text_box(slide, Inches(1.5), Inches(4.2), Inches(10), Inches(0.6),
                 "사용자 매뉴얼", font_size=28, color=TEXT_PRIMARY, bold=True)
    add_text_box(slide, Inches(1.5), Inches(5.0), Inches(10), Inches(0.5),
                 "전체 기능 안내서  |  2026년 6월", font_size=16, color=TEXT_SECONDARY)

    # bottom bar
    add_accent_bar(slide, Inches(0), Inches(7.38), SLIDE_W, Inches(0.06), ACCENT_BLUE)
    add_accent_bar(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), ACCENT_CYAN)


def make_toc_slide(sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    add_accent_bar(slide, Inches(0), Inches(0), Inches(0.08), SLIDE_H, ACCENT_BLUE)
    add_text_box(slide, Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
                 "목차 (Table of Contents)", font_size=30, color=TEXT_PRIMARY, bold=True)
    add_accent_bar(slide, Inches(0.6), Inches(1.15), Inches(2.5), Inches(0.04), ACCENT_CYAN)

    col1 = sections[:6]
    col2 = sections[6:]

    for i, sec in enumerate(col1):
        y = 1.6 + i * 0.8
        color = SECTION_COLORS[i] if i < len(SECTION_COLORS) else ACCENT_BLUE
        add_accent_bar(slide, Inches(0.6), Inches(y), Inches(0.06), Inches(0.5), color)
        add_text_box(slide, Inches(0.9), Inches(y), Inches(5.5), Inches(0.35),
                     sec, font_size=16, color=TEXT_PRIMARY)

    for i, sec in enumerate(col2):
        y = 1.6 + i * 0.8
        idx = i + 6
        color = SECTION_COLORS[idx] if idx < len(SECTION_COLORS) else ACCENT_BLUE
        add_accent_bar(slide, Inches(7.0), Inches(y), Inches(0.06), Inches(0.5), color)
        add_text_box(slide, Inches(7.3), Inches(y), Inches(5.5), Inches(0.35),
                     sec, font_size=16, color=TEXT_PRIMARY)


def make_section_slide(title, subtitle="", section_idx=0):
    global current_section_idx
    current_section_idx = section_idx

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    color = SECTION_COLORS[section_idx] if section_idx < len(SECTION_COLORS) else ACCENT_BLUE

    # large accent bar on left
    add_accent_bar(slide, Inches(0), Inches(0), Inches(0.12), SLIDE_H, color)

    # section number circle
    num = title.split(".")[0].strip() if "." in title else str(section_idx + 1)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), Inches(2.2), Inches(1.2), Inches(1.2))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(36)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "맑은 고딕"
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(12)

    # title text
    clean_title = title.split(".", 1)[-1].strip() if "." in title else title
    add_text_box(slide, Inches(3.2), Inches(2.0), Inches(9), Inches(1.0),
                 clean_title, font_size=40, color=TEXT_PRIMARY, bold=True)

    if subtitle:
        add_text_box(slide, Inches(3.2), Inches(3.2), Inches(9), Inches(0.6),
                     subtitle, font_size=18, color=TEXT_SECONDARY)

    add_accent_bar(slide, Inches(3.2), Inches(4.0), Inches(4), Inches(0.04), color)

    # bottom bar
    add_accent_bar(slide, Inches(0), Inches(7.38), SLIDE_W, Inches(0.06), color)


def make_content_slide(title, bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    color = SECTION_COLORS[current_section_idx] if current_section_idx < len(SECTION_COLORS) else ACCENT_BLUE

    # top accent
    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.05), color)

    # title
    add_text_box(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                 title, font_size=26, color=TEXT_PRIMARY, bold=True)
    add_accent_bar(slide, Inches(0.6), Inches(1.0), Inches(2), Inches(0.035), color)

    # content area with card background
    card = add_shape(slide, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.7),
                     fill_color=CARD_BG, border_color=CARD_BORDER, border_width=Pt(1))

    add_bullet_text(slide, Inches(0.9), Inches(1.5), Inches(11.5), Inches(5.3),
                    bullets, font_size=16, line_spacing=1.4)

    # page indicator
    slide_num = len(prs.slides)
    add_text_box(slide, Inches(12), Inches(7.1), Inches(1), Inches(0.3),
                 str(slide_num), font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)

    if note:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = note


def make_two_col_slide(title, left_title, left_bullets, right_title, right_bullets, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    color = SECTION_COLORS[current_section_idx] if current_section_idx < len(SECTION_COLORS) else ACCENT_BLUE

    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.05), color)

    add_text_box(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7),
                 title, font_size=26, color=TEXT_PRIMARY, bold=True)
    add_accent_bar(slide, Inches(0.6), Inches(1.0), Inches(2), Inches(0.035), color)

    # left card
    add_shape(slide, Inches(0.5), Inches(1.3), Inches(5.9), Inches(5.7),
              fill_color=CARD_BG, border_color=CARD_BORDER, border_width=Pt(1))
    if left_title:
        add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5.3), Inches(0.5),
                     left_title, font_size=16, color=color, bold=True)
        add_bullet_text(slide, Inches(0.8), Inches(1.9), Inches(5.3), Inches(4.9),
                        left_bullets, font_size=14, line_spacing=1.35)
    else:
        add_bullet_text(slide, Inches(0.8), Inches(1.5), Inches(5.3), Inches(5.3),
                        left_bullets, font_size=14, line_spacing=1.35)

    # right card
    add_shape(slide, Inches(6.9), Inches(1.3), Inches(5.9), Inches(5.7),
              fill_color=CARD_BG, border_color=CARD_BORDER, border_width=Pt(1))
    if right_title:
        add_text_box(slide, Inches(7.2), Inches(1.4), Inches(5.3), Inches(0.5),
                     right_title, font_size=16, color=color, bold=True)
        add_bullet_text(slide, Inches(7.2), Inches(1.9), Inches(5.3), Inches(4.9),
                        right_bullets, font_size=14, line_spacing=1.35)
    else:
        add_bullet_text(slide, Inches(7.2), Inches(1.5), Inches(5.3), Inches(5.3),
                        right_bullets, font_size=14, line_spacing=1.35)

    slide_num = len(prs.slides)
    add_text_box(slide, Inches(12), Inches(7.1), Inches(1), Inches(0.3),
                 str(slide_num), font_size=10, color=TEXT_MUTED, alignment=PP_ALIGN.RIGHT)

    if note:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = note


def make_ending_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    add_accent_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), ACCENT_BLUE)
    add_accent_bar(slide, Inches(0), Inches(0.06), SLIDE_W, Inches(0.03), ACCENT_CYAN)

    add_text_box(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(1),
                 "Thank You", font_size=48, color=ACCENT_BLUE, bold=True)
    add_text_box(slide, Inches(1.5), Inches(3.4), Inches(10), Inches(0.6),
                 "DC Express — Data Center Infrastructure Management", font_size=20, color=TEXT_MUTED)

    add_accent_bar(slide, Inches(1.5), Inches(4.3), Inches(3), Inches(0.04), ACCENT_CYAN)

    add_text_box(slide, Inches(1.5), Inches(4.8), Inches(10), Inches(0.5),
                 "문의사항이 있으시면 언제든지 연락 주세요.", font_size=16, color=TEXT_SECONDARY)

    add_accent_bar(slide, Inches(0), Inches(7.38), SLIDE_W, Inches(0.06), ACCENT_BLUE)
    add_accent_bar(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06), ACCENT_CYAN)


def process_slide(slide_data, section_idx=0):
    global current_section_idx
    layout = slide_data.get("layout", "content")

    if layout == "section":
        current_section_idx = section_idx
        make_section_slide(
            slide_data.get("title", ""),
            slide_data.get("subtitle", ""),
            section_idx
        )
    elif layout == "two_col":
        current_section_idx = section_idx
        make_two_col_slide(
            slide_data.get("title", ""),
            slide_data.get("left_title", ""),
            slide_data.get("left_bullets", []),
            slide_data.get("right_title", ""),
            slide_data.get("right_bullets", []),
            slide_data.get("note", "")
        )
    else:
        current_section_idx = section_idx
        make_content_slide(
            slide_data.get("title", ""),
            slide_data.get("bullets", []),
            slide_data.get("note", "")
        )


# ============================================================
# SLIDE DATA — will be populated by agents
# ============================================================

# --- Section 1-2: Dashboard + Server Monitoring ---
SECTION_1_2 = [
    {
        "layout": "section",
        "title": "1. 대시보드 (Dashboard)",
        "subtitle": "인프라 전체 현황을 한눈에 파악하는 통합 모니터링 화면",
    },
    {
        "layout": "content",
        "title": "대시보드 개요",
        "bullets": [
            "브라우저에서 DC Express 주소(http://서버IP:3000)로 접속하면 대시보드가 기본 화면으로 표시",
            "좌측 사이드바 메뉴의 Dashboard 항목을 클릭해서도 이동 가능",
            "화면 구성: 클러스터 필터 탭 + 실시간 메트릭 카드(8개) + Fleet Top 5 + 서버 상태 차트 + 사이드바 위젯",
            "Prometheus 데이터를 30초 간격으로 자동 갱신하여 실시간 현황 반영",
            "Prometheus 연결 장애 시에도 에러 바운더리로 UI 깨짐 방지",
        ],
    },
    {
        "layout": "content",
        "title": "클러스터 필터",
        "bullets": [
            "대시보드 상단에 All / Lab-1 / Lab-3 세 개의 탭 버튼 제공",
            "All — 전체 장비 통합 현황 (기본값)",
            "Lab-1 — 10.144.38.x 대역 장비만 표시",
            "Lab-3 — 10.144.131.x 대역 장비만 표시",
            "탭 전환 시 모든 메트릭 카드, 차트, Fleet Top 5가 해당 클러스터 기준으로 즉시 갱신",
        ],
    },
    {
        "layout": "two_col",
        "title": "실시간 메트릭 카드 (1/2) — 상단 4개",
        "left_title": "좌측",
        "left_bullets": [
            "Average CPU — 전체 서버 평균 CPU 사용률(%)",
            "• 스파크라인(30분 추이 그래프) 포함",
            "Average Memory — 전체 서버 평균 메모리 사용률(%)",
            "• 스파크라인 포함",
        ],
        "right_title": "우측",
        "right_bullets": [
            "Temperature — CPU 온도 + DIMM 온도(°C)",
            "• 두 값 병렬 표시",
            "Total Power — 전체 소비 전력(W/kW)",
            "• Intel PCM Package_Joules 합산, 자동 단위 변환",
        ],
    },
    {
        "layout": "two_col",
        "title": "실시간 메트릭 카드 (2/2) — 하단 4개",
        "left_title": "좌측",
        "left_bullets": [
            "Nodes — 정상(UP) / 비정상(DOWN) 노드 수",
            "• 전체 정상: 초록 아이콘 / 비정상 존재: 빨강 아이콘",
            "Average Uptime — 평균 서버 가동 시간(일)",
        ],
        "right_title": "우측",
        "right_bullets": [
            "Network In — 수신(RX) 트래픽 + TX 값 병기",
            "• 스파크라인 포함",
            "PUE — 전력 효율 지수 (아크 게이지)",
            "• 1.4 미만 Good / 1.4~1.6 Average / 1.6 이상 Poor",
        ],
    },
    {
        "layout": "two_col",
        "title": "Fleet Top 5 + 서버 상태",
        "left_title": "Fleet Top 5",
        "left_bullets": [
            "Fleet Top 5 — CPU: 사용률 상위 5대 막대 그래프",
            "• 80% 초과 빨강 / 60% 초과 주황 / 이하 시안",
            "Fleet Top 5 — Memory: 사용률 상위 5대 막대 그래프",
            "• 80% 초과 빨강 / 60% 초과 주황 / 이하 초록",
            "서버 이름 클릭 시 해당 서버 상세 페이지로 즉시 이동",
        ],
        "right_title": "서버 상태",
        "right_bullets": [
            "Server Status 도넛 차트",
            "• Active(초록) / Maintenance(주황) / Failed(빨강)",
            "• 중앙에 총 서버 수 표시",
            "Platform Status — 플랫폼별 카드",
            "• SPR, GNR-AP, GNR-SP, SRF, Ampere, EMR 등",
            "• 플랫폼별 가동률과 오프라인 수 표시",
        ],
    },
    {
        "layout": "content",
        "title": "사이드바 위젯 + 요약 카드 + Quick Links",
        "bullets": [
            "Active Alerts — 현재 발생 중인 알림 상위 5건 (심각도별 색상: 빨강/주황/파랑)",
            "• History 클릭 시 알림 관리 페이지(/alerts)로 이동",
            "Expiring Soon — 30일 이내 만료 예정 자산, D-day 카운트다운 표시",
            "Total Equipment 카드 — Active/Maintenance/Failed 장비 수",
            "• 마우스 오버 시 해당 장비 목록(최대 10개) 팝오버 표시",
            "Infrastructure 카드 — 총 랙 수와 룸 수 표시",
            "Quick Links — Servers, Infrastructure, Alerts, Discovery 바로가기",
        ],
    },
    {
        "layout": "section",
        "title": "2. 서버 모니터링 (Server Monitoring)",
        "subtitle": "개별 서버의 상세 메트릭과 비교 분석 기능",
    },
    {
        "layout": "content",
        "title": "서버 목록 — 리스트 뷰",
        "bullets": [
            "사이드바 Servers 메뉴 또는 /servers 경로로 접속",
            "통합 검색 — hostname, IP, BMC IP, 모델명, 제조사, CPU 정보, Room/Rack 이름",
            "5개 필터 — Model / Status / Room / Power / 초기화 버튼(활성 필터 수 표시)",
            "정렬 가능 컬럼: Hostname, IP, Status, System Model (헤더 클릭으로 전환)",
            "테이블 컬럼: Hostname, IP, BMC IP, Status, Power, CPU, System Model, Room, Rack/U",
        ],
    },
    {
        "layout": "two_col",
        "title": "서버 목록 — Power State + 트윈 뷰",
        "left_title": "Power State 배지",
        "left_bullets": [
            "Running (초록, 점 깜빡임)",
            "• CPU 사용률 5% 초과 — 워크로드 실행 중",
            "Idle (노랑)",
            "• node-exporter 응답 중이나 CPU 5% 이하",
            "OFF (회색)",
            "• node-exporter 무응답 — 서버 꺼짐/네트워크 단절",
            "30초마다 자동 갱신",
        ],
        "right_title": "트윈 뷰 (Twin View)",
        "right_bullets": [
            "List / Twin 토글로 뷰 전환",
            "3단계 드릴다운 구조:",
            "• 데이터센터 레벨 — SVG 평면도, 오버레이 지원",
            "• Room 레벨 — 랙 배치도, 사용률 바",
            "• Rack 레벨 — U별 장비, 클릭→상세 이동",
        ],
    },
    {
        "layout": "two_col",
        "title": "서버 상세 — 개요 및 시간 범위",
        "left_title": "상단 정보",
        "left_bullets": [
            "IP, BMC IP, Status, Model, Location(Room/Rack/U)",
            "CPUs — 개수 x 모델 + 아키텍처",
            "OS 정보 / Memory 총 용량 + DIMM 현황",
            "BMC 서버: Refresh HW + Power Console 표시",
        ],
        "right_title": "시간 범위 선택",
        "right_bullets": [
            "15m (15초 간격)",
            "1h (30초 간격) — 기본값",
            "6h (2분 간격)",
            "24h (5분 간격)",
            "7d (30분 간격)",
            "변경 시 모든 차트 동시 갱신",
        ],
    },
    {
        "layout": "two_col",
        "title": "서버 상세 — CPU 메트릭 (7개 차트)",
        "left_title": "기본 CPU 차트",
        "left_bullets": [
            "CPU Usage — 전체 사용률 추이(%)",
            "CPU Load Average — 1m/5m/15m + 코어 수 기준선",
            "CPU Mode Breakdown — user/system/iowait/steal",
            "CFS Throttled — 컨테이너 스로틀링(sec/s)",
        ],
        "right_title": "고급 CPU 차트",
        "right_bullets": [
            "PCM IPC — 코어당 명령어 처리량 (Intel PCM 전용)",
            "PCM Cache Hit Rate — L2/L3 캐시 히트율(%)",
            "Core Heatmap — 코어별 사용률 격자",
            "• 초록(낮음) ~ 빨강(높음) 색상 변화",
            "• Pod별 CPU 바 차트 대체 모드",
        ],
    },
    {
        "layout": "two_col",
        "title": "서버 상세 — Memory / Disk / Network / Hardware",
        "left_title": "Memory + Disk",
        "left_bullets": [
            "Memory Usage (%) + Swap 추이",
            "Memory Absolute — Used / Cache+Buffer",
            "PCM Memory Bandwidth — 읽기/쓰기 대역폭",
            "Disk I/O — Throughput, IOPS, Latency",
            "Disk Usage (%) + Filesystem Breakdown 표",
        ],
        "right_title": "Network + Hardware",
        "right_bullets": [
            "Bandwidth (RX/TX), Errors, TCP Connections",
            "TCP Retransmits, NIC Inventory(UP/DOWN)",
            "Temperature — hwmon + Inlet/Exhaust/Socket 통합",
            "Power — RAPL: Package/DRAM/PP0 (Watts)",
            "Fan Speed — RPM",
        ],
    },
    {
        "layout": "content",
        "title": "서버 비교 기능",
        "bullets": [
            "서버 목록 우측 상단 Compare 버튼 → /servers/compare 이동",
            "hostname, IP, Room, Rack으로 검색 후 클릭하여 선택 (최대 4대)",
            "선택 서버별 색상 점 표시 (파랑, 초록, 주황, 보라 순서)",
            "최소 2대 선택 시 6개 오버레이 차트가 2열 격자로 표시:",
            "• CPU Usage(%) / Memory Usage(%) / Load Average 1m",
            "• Disk I/O Read(Bytes/s) / Network RX(Bytes/s) / Temperature(°C)",
            "시간 범위(15m/1h/6h/24h/7d) 변경 시 6개 차트 동시 갱신",
        ],
    },
]

# --- Section 3-4: Rack View + Digital Twin ---
SECTION_3_4 = [
    {
        "layout": "section",
        "title": "3. 랙 뷰 (Rack View)",
        "subtitle": "물리적 랙 배치와 장비 구성을 한눈에 파악",
    },
    {
        "layout": "content",
        "title": "랙 페이지 개요",
        "bullets": [
            "사이드바 Racks 메뉴 또는 /racks 경로로 접속",
            "페이지 상단에 4개의 요약 카드 표시:",
            "• Total Racks — 등록된 전체 랙 수",
            "• Rooms — 방(Room) 수",
            "• Equipment — 랙에 배치된 전체 장비 수",
            "• U Utilization — 전체 U 공간 사용률 (사용 중 U / 전체 U)",
            "요약 카드 아래에 Room별로 그룹화된 랙 목록이 표시됨",
        ],
    },
    {
        "layout": "content",
        "title": "Room별 랙 그룹 + 랙 카드 정보",
        "bullets": [
            "랙은 소속 Room 단위로 그룹화 — Room 이름 옆에 랙 수 배지 표시",
            "각 랙 카드 요약 정보:",
            "• 랙 이름 및 Row 라벨",
            "• U 사용률 프로그레스 바 (85% 초과: 빨강, 60% 초과: 주황, 이하: 초록)",
            "• 장비 수, 활성(Active) 장비 수, 장애(Failed) 장비 수",
            "• 최대 전력 용량 (설정된 경우 표시)",
            "랙 카드 클릭 시 Rack Elevation이 펼쳐짐 (아코디언 패턴)",
        ],
    },
    {
        "layout": "two_col",
        "title": "Rack Elevation 시각화",
        "left_title": "인터랙티브 Rack Elevation",
        "left_bullets": [
            "U1부터 최상단까지 세로 슬롯 다이어그램",
            "장비 배치 슬롯은 상태별 색상으로 표시:",
            "• ACTIVE:초록 / MAINTENANCE:보라 / REPAIR:주황",
            "• FAILED:빨강 / PLANNED:파랑 / INSTALLED:청록",
            "장비 호버 시 팝오버(U 위치, 모델, IP)",
            "빈 슬롯은 'empty' 표시",
        ],
        "right_title": "장비 테이블 + 조작",
        "right_bullets": [
            "U 위치, Hostname, IP, 상태, 타입, 모델 표시",
            "Hostname 클릭 → 서버 상세 페이지 이동",
            "위/아래 화살표로 장비 1U씩 이동",
            "X 버튼으로 랙에서 장비 제거(배치 해제)",
            "한 번에 하나의 랙만 펼칠 수 있음",
        ],
    },
    {
        "layout": "two_col",
        "title": "온도 히트맵 + 드래그앤드롭 배치",
        "left_title": "온도 히트맵",
        "left_bullets": [
            "'온도 ON/OFF' 버튼으로 히트맵 모드 토글",
            "온도 색상 6단계:",
            "• <35°C 하늘색 → 35~45 초록 → 45~55 노랑",
            "• 55~65 주황 → 65~75 오렌지 → 75°C+ 빨강",
            "장비 우측에 실시간 온도 수치 표시",
            "Prometheus에서 30초 갱신",
        ],
        "right_title": "드래그앤드롭",
        "right_bullets": [
            "장비 좌측 그립 → 원하는 빈 슬롯에 드롭",
            "'여기에 놓기' 파란색 안내 표시",
            "U 충돌 방지 내장 — 중복 배치 불가",
            "미배치 장비 추가 패널:",
            "• 하단 '장비 추가' 버튼 → 장비 선택 → U 지정",
            "• 미지정 시 첫 번째 빈 슬롯 자동 배치",
        ],
    },
    {
        "layout": "content",
        "title": "랙 관리 페이지 (/racks/manage)",
        "bullets": [
            "랙 페이지 상단 '관리' 버튼 → Admin/Operator 권한 필요",
            "[Room/Rack 관리 탭]",
            "• Room CRUD — 이름, 설명, 정렬 순서(sortOrder)",
            "• Rack CRUD — 이름, Row, 전체 U 수(기본 42U), 최대 전력(W), 평면도 좌표(X/Y)",
            "• 장비가 배치된 랙은 삭제 전 장비 이동/해제 필요",
            "[일괄 배치 탭]",
            "• 다수의 장비를 한 번에 특정 랙에 일괄 배치",
        ],
    },
    {
        "layout": "section",
        "title": "4. 디지털 트윈 (Digital Twin)",
        "subtitle": "데이터센터 물리 구조를 SVG 기반 평면도로 시각화",
    },
    {
        "layout": "content",
        "title": "데이터센터 평면도 구성",
        "bullets": [
            "건물 전체를 1400x650 크기의 SVG 캔버스로 표현",
            "Lab-3 (상단) — 보라색 강조, 서버 랙 + 스위치, 냉각기, PDU, K8s 마스터",
            "Lab-2 (하단 좌측) — 녹색 강조, 확장 예정 공간 (AVAILABLE SPACE)",
            "Lab-1 (하단 우측) — 파란색 강조, 서버 랙 + 인프라 요소",
            "각 방 요약: 랙 수, 서버 수, Active 수, U 사용률 바, 상태 배지",
            "방과 방 사이에 문(DOOR) 마커로 동선 표시",
        ],
    },
    {
        "layout": "content",
        "title": "오버레이 모드 (3가지)",
        "bullets": [
            "평면도 상단 3개 토글 버튼 — 한 번에 하나만 활성화",
            "[온도 Temp] Prometheus hwmon 데이터 30초 갱신, 랙별 평균 온도 색상",
            "• 30°C 이하 초록 → 45 라임 → 55 노랑 → 65 주황 → 75+ 빨강",
            "[용량 Capacity] 랙별 U 사용률 색상",
            "• 60% 미만 초록 / 60~85% 주황 / 85% 초과 빨강",
            "[에어플로우 Airflow] 냉기 흐름 방향 화살표 + COLD→HOT 라벨",
            "줌/팬: 마우스 드래그=팬, +/-/리셋 버튼, 0.5x~4x 범위",
        ],
    },
    {
        "layout": "content",
        "title": "랙 호버 팝오버 + 클릭 이동",
        "bullets": [
            "랙에 마우스 호버 시 HTML 오버레이 팝오버 표시",
            "• 랙 이름 + '클릭하여 이동' 안내",
            "• 서버 수: 활성/전체 (장애 시 빨간 배지)",
            "• U 사용률: 퍼센트 + 사용/전체 U + 프로그레스 바",
            "• 평균 온도: 색상 구간 표시",
            "• 서버 호스트명 목록: 최대 6개, 초과 시 '+N' 배지",
            "랙 클릭 → /racks?highlight=랙ID → 해당 랙 자동 펼침 + 스크롤",
        ],
    },
    {
        "layout": "two_col",
        "title": "편집 모드 (관리자용)",
        "left_title": "기본 조작",
        "left_bullets": [
            "연필 아이콘 → 편집 모드 진입",
            "좌클릭 드래그 → 크기 조절 (W/H)",
            "우클릭 드래그 → 위치 이동 (방 경계 내)",
            "+ 버튼 → 요소 추가 (COOLING/PDU/SWITCH 등)",
            "x 버튼 → 요소 삭제",
            "변경 사항 API로 즉시 저장",
        ],
        "right_title": "고급 기능",
        "right_bullets": [
            "크기 통일: 우클릭 메뉴 → 동일 타입 일괄 적용",
            "방향 전환: 냉각기(상/우/하/좌), 문(가로/세로)",
            "다중 선택: 좌클릭으로 2개+ 선택",
            "• 가로 정렬: Y 좌표 평균값 통일",
            "• 세로 정렬: X 좌표 평균값 통일",
            "방 이동/크기 조절 (최소 100x80)",
        ],
    },
]

# --- Section 5-7: BMC + K8s + Workloads ---
SECTION_5_7 = [
    {
        "layout": "section",
        "title": "5. BMC/Redfish 제어",
        "subtitle": "제조사 독립적 서버 원격 관리",
    },
    {
        "layout": "content",
        "title": "BMC 개요 — Redfish 표준 기반 원격 관리",
        "bullets": [
            "BMC(Baseboard Management Controller): 서버 메인보드 내장 독립 관리 프로세서",
            "• OS와 별개 동작 — 서버 꺼진 상태에서도 전원 제어, 센서 읽기 가능",
            "Redfish: DMTF 정의 RESTful 서버 관리 표준 API",
            "• Intel, Dell, HPE, Supermicro 등 제조사 독립적 통합 제어",
            "DC Express는 Redfish API를 통해 BMC와 직접 통신",
            "• 자체 서명 인증서 환경 지원 (내부 네트워크 기본 설정)",
        ],
    },
    {
        "layout": "content",
        "title": "BMC IP 설정 — 3가지 방법",
        "bullets": [
            "방법 1: 장비 등록/수정 시 직접 입력",
            "• /infrastructure/new 또는 /infrastructure/장비ID/edit의 BMC IP 필드",
            "방법 2: BMC 관리 페이지에서 인라인 편집",
            "• 설정 > BMC 관리 > BMC IP 매핑 탭 — 연필 아이콘 클릭 → 즉석 편집",
            "• 필터: 전체 / BMC 설정됨 / BMC 미설정",
            "방법 3: 자동 유도 (일괄 자동 설정)",
            "• Host IP 마지막 옥텟 → 192.168.10.{옥텟} 자동 생성",
            "• '미설정 서버 자동 설정' 버튼으로 일괄 적용",
        ],
    },
    {
        "layout": "two_col",
        "title": "BMC 자격증명 + 프록시",
        "left_title": "기본 자격증명",
        "left_bullets": [
            "환경 변수 설정:",
            "• BMC_USERNAME / BMC_PASSWORD",
            "• .env 파일에 설정, 모든 서버 공통 적용",
            "• 미설정 시 BMC 기능 전체 비활성화",
            "장비별 계정 오버라이드 가능",
        ],
        "right_title": "BMC 프록시 (원격 사이트)",
        "right_bullets": [
            "네트워크 분리된 원격 사이트(Lab-3 등) BMC 접근",
            "Room 설정에서 BMC 프록시 URL 지정",
            "• 예: http://10.144.131.100:8443",
            "Room 단위 설정 → 같은 방 모든 장비 자동 적용",
            "직접 접근 대신 프록시 서버 경유",
        ],
    },
    {
        "layout": "content",
        "title": "BMC 센서 모니터링",
        "bullets": [
            "서버 상세 페이지에 BMC Sensors (Thermal & Power) 카드 자동 표시",
            "전력 소비: 현재 소비 전력(W) + 전체 용량 대비 비율 프로그레스바",
            "전원 공급 장치(PSU): PSU별 이름, 모델, 용량(W), 상태 — 카드 그리드",
            "온도 센서: 현재 온도(°C) + 임계값 대비 바 차트 (녹/황/적)",
            "• 임계값(Critical) 85% 이상 황색, 초과 시 적색",
            "팬 센서: 팬별 회전 속도(RPM) 숫자 그리드",
            "Refresh 버튼으로 최신 데이터 재조회 + 조회 시각 표시",
        ],
    },
    {
        "layout": "content",
        "title": "전원 제어 (Power Control) — 5가지 액션",
        "bullets": [
            "서버 상세 Power & Console 카드에서 전원 상태 확인 및 제어",
            "Graceful Restart — OS 정상 재부팅 (위험도: 보통)",
            "Force Restart — BMC 강제 재시작, 데이터 유실 가능 (위험도: 주의)",
            "Graceful Shutdown — OS 정상 종료 (위험도: 주의)",
            "Force Off — 전원 즉시 차단 = 전원 코드 제거와 동일 (위험도: 위험)",
            "Power On — 시스템 전원 켜기 (위험도: 보통)",
        ],
    },
    {
        "layout": "content",
        "title": "전원 제어 — 안전 장치",
        "bullets": [
            "호스트명 직접 타이핑: 대상 서버 호스트명을 정확히 입력해야 실행 (실수 방지)",
            "사유 입력 필수: 최소 3자 이상의 전원 제어 사유 기록",
            "티켓 참조 (선택): 관련 작업 티켓 번호 연결",
            "역할 기반 접근 제어: ADMIN/OPERATOR만 제어, VIEWER는 조회만",
            "감사 로그 자동 기록: 모든 전원 제어 이력 자동 저장",
            "일괄 전원 제어: 설정>BMC 관리>일괄 작업 탭에서 다수 서버 동시 제어",
        ],
    },
    {
        "layout": "content",
        "title": "하드웨어 자동 수집 (Refresh HW)",
        "bullets": [
            "Redfish API로 BMC에서 HW 정보 자동 수집 → DB 저장",
            "수집 항목:",
            "• CPU — 소켓별 모델, 코어/스레드, TDP, 아키텍처",
            "• Memory — 슬롯별 용량, 타입(DDR4/DDR5), 제조사, 시리얼, 속도",
            "• NIC — 포트명, MAC, 링크 속도, IPv4",
            "• 시스템 — 제조사, 모델명, S/N, BIOS 버전",
            "일괄 HW 갱신: 다수 서버 동시 수집 → 수동 입력 없이 자동 완성",
        ],
    },
    {
        "layout": "section",
        "title": "6. 쿠버네티스(K8s) 연동",
        "subtitle": "Prometheus 기반 K8s 메트릭 통합",
    },
    {
        "layout": "content",
        "title": "K8s 메트릭 수집 구조",
        "bullets": [
            "별도 K8s API 직접 연결 없이 기존 Prometheus 인프라 활용",
            "kube-state-metrics (1대): 노드/파드 상태, 용량 정보",
            "• kube_pod_info, kube_node_status_capacity, kube_pod_status_phase",
            "kubernetes-cadvisor (24대): 컨테이너별 CPU/메모리/디스크/네트워크",
            "kubelet (각 노드): 파드 실행 현황 — kubelet_running_pods",
            "데이터 흐름: K8s 클러스터 → Prometheus → DC Express API → 웹 UI",
        ],
    },
    {
        "layout": "two_col",
        "title": "Prometheus Discovery + 서버 K8s 정보",
        "left_title": "Prometheus Discovery",
        "left_bullets": [
            "설정 > Prometheus Discovery에서 접근",
            "동기화 실행: /api/v1/targets 호출 → 자동 수집",
            "필터: Job별 / UP·DOWN / 연결·미연결 / 텍스트 검색",
            "미연결 타겟 '등록' → 장비 자동 생성",
            "CPU 코어 수, 메모리 용량 자동 감지",
        ],
        "right_title": "Node Resources 카드",
        "right_bullets": [
            "서버 상세 페이지에 자동 표시",
            "CPU/Memory/Disk/Pods 게이지",
            "• 70% 이상 노란색, 90% 이상 빨간색",
            "Running Pods 목록",
            "• namespace/pod이름 형식",
            "• kube_pod_info 메트릭 기반",
        ],
    },
    {
        "layout": "section",
        "title": "7. 워크로드 관리",
        "subtitle": "Workload Management",
    },
    {
        "layout": "content",
        "title": "Active 탭 — 실시간 워크로드 모니터링",
        "bullets": [
            "Prometheus에서 실시간 파드 상태 조회 → 네임스페이스별 카드 표시",
            "각 카드: Health 상태 표시등, 파드/노드 수, 시작일/경과 시간",
            "Health 상태 5종류:",
            "• Running(녹) — 정상 / Pending(노랑) — 대기 / Warning(주황) — 경고",
            "• Error(빨강) — CrashLoopBackOff 등 / Completed(회색) — 완료",
            "시스템 NS(kube-system, monitoring 등) 자동 제외",
            "카드 클릭 → 상세 페이지 (파드 목록, 프로젝트, 테스트, 태스크, 메모)",
        ],
    },
    {
        "layout": "content",
        "title": "History 탭 — 달력 타임라인",
        "bullets": [
            "달력 타임라인 형태로 모든 워크로드 이력 시각화",
            "네임스페이스 필터: 체크박스 선택, LIVE 표시, 고유 색상 구분",
            "월 단위 달력에 프로젝트 기간을 색상 바로 표시",
            "프로젝트 상세 팝업 (바 클릭): Overview, Phases, Test Results, Tasks, Notes",
            "전체 워크로드 리스트: 상태 5종 (IN_PROGRESS/COMPLETED/CANCELLED/ON_HOLD/PLANNED)",
        ],
    },
    {
        "layout": "content",
        "title": "대시보드 연동 — Active Workloads 카드",
        "bullets": [
            "메인 대시보드에 Active Workloads 카드 — 워크로드 실시간 요약",
            "NS별 Health 상태, 파드/노드 수, 시작일, 경과 시간",
            "Warning 파드 시 사유(예: CrashLoopBackOff)와 파드명 표시",
            "노드 온도 바: 50°C↓ 녹 / 65°C↑ 주황 / 80°C↑ 빨강",
            "클러스터 필터 연동: All / Lab-1 / Lab-3",
            "카드 클릭 → /workloads/NS명 상세 이동, 30초 자동 갱신",
        ],
    },
]

# --- Section 8-12: Infrastructure + Alerts + Settings + Others ---
SECTION_8_12 = [
    {
        "layout": "section",
        "title": "8. 인프라 관리",
        "subtitle": "Infrastructure Management",
    },
    {
        "layout": "content",
        "title": "장비 목록 및 검색",
        "bullets": [
            "전체 장비 테이블 뷰: hostname, IP, 모델, 상태, 위치, 메모리, CPU 컬럼",
            "플랫폼별 필터 칩: GNR-AP, GNR-SP, SPR, SRF, Ampere 등 자동 매핑",
            "텍스트 검색: hostname, IP, 모델명으로 실시간 필터링",
            "상태 요약 카드: ACTIVE, MAINTENANCE, RETIRED 등 상태별 장비 수",
            "장비 클릭 시 상세 이동, ADMIN/OPERATOR에게 등록/Bulk Import 버튼 노출",
        ],
    },
    {
        "layout": "two_col",
        "title": "장비 등록 / 수정",
        "left_title": "개별 등록",
        "left_bullets": [
            "장비 타입 9종: SERVER, SWITCH, ROUTER 등",
            "필수 입력: hostname, IP, 모델, 타입",
            "CPU 소켓별 상세 (코어/스레드/TDP/아키텍처)",
            "메모리 슬롯별 입력 (DDR3~HBM3)",
            "랙 위치: Room/Rack/U position 선택",
        ],
        "right_title": "대량 등록 + 라이프사이클",
        "right_bullets": [
            "CSV 대량 등록: 업로드 → 미리보기 → Zod 검증",
            "라이프사이클 관리:",
            "• REGISTERED → ACTIVE → MAINTENANCE",
            "• → RETIRED → DECOMMISSIONED",
            "상태 변경 이력 자동 기록",
        ],
    },
    {
        "layout": "content",
        "title": "장비 상세 + 부가 기능",
        "bullets": [
            "기본 정보: hostname, IP, 모델, S/N, 상태, Power State 실시간",
            "CPU/Memory/NIC 사양 상세, DIMM 슬롯 시각화 (/memory 별도 페이지)",
            "BMC 연동: HW 정보 갱신 + 콘솔 링크 + 원격 전원 제어",
            "변경 이력 타임라인: 수정/전원/상태/유지보수 이력 시간순 표시",
            "사용자 할당: 장비별 담당자 할당/반납 및 이력 추적",
            "펌웨어 관리: 모델별 BIOS/펌웨어 버전 비교, 오래된 버전 자동 식별",
        ],
    },
    {
        "layout": "section",
        "title": "9. 알림 관리",
        "subtitle": "Alert Management",
    },
    {
        "layout": "content",
        "title": "알림 규칙 + 이력 관리",
        "bullets": [
            "PromQL 기반 알림 규칙 CRUD (Zod 유효성 검증)",
            "심각도 3단계: CRITICAL(빨강), WARNING(주황), INFO(파랑)",
            "카테고리: temperature, cpu, memory, disk, network, power, hardware",
            "프리셋 5종: High CPU(>90%), High Memory(>90%), Disk Low(<10%), High Temp(>80°C), Node Down",
            "알림 이력: 상태 카드(Firing/Acknowledged/Resolved), 날짜별 Accordion, 카테고리 필터",
            "워크플로우: FIRING → ACKNOWLEDGED → RESOLVED (역할 기반 권한)",
        ],
    },
    {
        "layout": "content",
        "title": "알림 고급 기능",
        "bullets": [
            "알림 수신 채널: Email, Slack, Teams, Webhook 4종 지원",
            "• 채널별 최소 심각도 설정 + 테스트 발송 + 활성/비활성 토글",
            "에스컬레이션 정책: 심각도별 N분 미확인 시 자동 통지",
            "• 예: CRITICAL 15분 미확인 → Ops Slack 전송",
            "유지보수 창 (Maintenance Windows):",
            "• 기간 + 범위별 알림 음소거 — muted 배지 표시",
            "해결 완료 알림 일괄 삭제 (ADMIN 전용)",
        ],
    },
    {
        "layout": "section",
        "title": "10. 설정",
        "subtitle": "Settings",
    },
    {
        "layout": "two_col",
        "title": "설정 페이지 + 사용자/역할 관리",
        "left_title": "설정 메뉴 (6개)",
        "left_bullets": [
            "사용자 관리 — 계정 CRUD, 역할 변경",
            "Prometheus Discovery — 타겟 동기화",
            "BMC 관리 — 일괄 HW 갱신, 전원 제어",
            "Prometheus 진단 — IP-only, 중복, 고아 탐지",
            "만기 관리 — 인증서/라이선스/보증 D-day",
            "시스템 정보 — 앱 버전, 환경, Prometheus URL",
        ],
        "right_title": "역할 기반 접근 제어",
        "right_bullets": [
            "Admin — 모든 기능 + 사용자/설정 관리",
            "Operator — 장비 CRUD + 알림 처리 + 전원 제어",
            "Viewer — 읽기 전용 (대시보드/모니터링)",
            "감사 로그 (/history):",
            "• 액션/엔티티/사용자/날짜별 필터",
            "• 변경 전후 diff + CSV 내보내기",
        ],
    },
    {
        "layout": "section",
        "title": "11. 용량 계획 / 리포트",
        "subtitle": "Capacity Planning & Reports",
    },
    {
        "layout": "two_col",
        "title": "용량 계획 + 리포트",
        "left_title": "Capacity Planning",
        "left_bullets": [
            "요약 카드: 장비 수, Rack Space, 총 메모리, TDP",
            "Room별 사용률 프로그레스 바 (4단계 색상)",
            "Rack별 상세 테이블 (장비, U 사용률, 전력)",
            "용량 예측: 월별 성장 추이 AreaChart",
            "• 선형회귀 12개월 전망",
            "• 랙 공간/전력 소진 예상일 카드",
        ],
        "right_title": "Reports",
        "right_bullets": [
            "Executive Summary: 장비, 메모리, Room/Rack, 알림",
            "Equipment Breakdown: 상태/타입/제조사별 분포",
            "Alert Statistics (30d): 카테고리별, Top Rules",
            "Recent Alerts: 최근 30일 알림 목록",
            "내보내기: PDF (A4, 한글 지원) + CSV + 브라우저 인쇄",
        ],
    },
    {
        "layout": "section",
        "title": "12. 기타 기능",
        "subtitle": "글로벌 검색, 다국어, 인증, 운영 자동화",
    },
    {
        "layout": "two_col",
        "title": "검색 + UX + 인증 + 운영",
        "left_title": "글로벌 검색 + UX",
        "left_bullets": [
            "Cmd+K 커맨드 팔레트 — 서버/Room/Rack/알림 통합 검색",
            "키보드 탐색 (화살표 + Enter)",
            "다국어 (한/영) — i18n 실시간 전환",
            "다크/라이트 모드 — CSS 변수 반전",
            "• localStorage 저장 + 무플래시 초기화",
        ],
        "right_title": "인증 + 운영 자동화",
        "right_bullets": [
            "NextAuth.js + JWT 전략",
            "역할 기반 접근 제어 (Admin/Operator/Viewer)",
            "API 20+ 라우트 Zod 입력 검증",
            "systemd 서비스 + Docker 배포",
            "DB 백업(매일) + 로그 로테이션 + 장애 복구",
        ],
    },
]


# ============================================================
# MAIN
# ============================================================

def main():
    # 1. Cover
    make_cover_slide()

    # 2. TOC
    toc_sections = [
        "1. 대시보드 (Dashboard)",
        "2. 서버 모니터링 (Server Monitoring)",
        "3. 랙 뷰 (Rack View)",
        "4. 디지털 트윈 (Digital Twin)",
        "5. BMC/Redfish 제어",
        "6. 쿠버네티스(K8s) 연동",
        "7. 워크로드 관리 (Workload Management)",
        "8. 인프라 관리 (Infrastructure)",
        "9. 알림 관리 (Alert Management)",
        "10. 설정 (Settings)",
        "11. 용량 계획 / 리포트",
        "12. 기타 기능",
    ]
    make_toc_slide(toc_sections)

    # 3. Slides
    section_map = {
        0: SECTION_1_2,   # section_idx 0 = Dashboard
        1: SECTION_1_2,   # section_idx 1 = Server Monitoring
        2: SECTION_3_4,   # 2 = Rack View
        3: SECTION_3_4,   # 3 = Digital Twin
        4: SECTION_5_7,   # 4 = BMC
        5: SECTION_5_7,   # 5 = K8s
        6: SECTION_5_7,   # 6 = Workloads
        7: SECTION_8_12,  # 7 = Infrastructure
        8: SECTION_8_12,  # 8 = Alerts
        9: SECTION_8_12,  # 9 = Settings
        10: SECTION_8_12, # 10 = Capacity
        11: SECTION_8_12, # 11 = Others
    }

    all_slides = []
    for slides_data in [SECTION_1_2, SECTION_3_4, SECTION_5_7, SECTION_8_12]:
        if isinstance(slides_data, list):
            all_slides.extend(slides_data)

    current_sec_idx = 0
    for sd in all_slides:
        if sd.get("layout") == "section":
            title = sd.get("title", "")
            for i, sec_name in enumerate(toc_sections):
                num = sec_name.split(".")[0].strip()
                if title.startswith(num + ".") or title.startswith(num + " "):
                    current_sec_idx = i
                    break
        process_slide(sd, current_sec_idx)

    # 4. Ending
    make_ending_slide()

    # Save
    output_path = os.path.join(os.path.dirname(__file__), "..", "docs", "DC_Express_User_Manual.pptx")
    output_path = os.path.abspath(output_path)
    prs.save(output_path)
    print(f"PowerPoint saved: {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
